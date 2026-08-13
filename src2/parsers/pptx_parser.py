"""
PPTX parser based on python-pptx and zip extraction.
"""

import os
import zipfile
import tempfile
import shutil
import xml.etree.ElementTree as ET
import re
from typing import Dict, Any, List

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .base_parser import BaseParser
from utils.image_filters import repeated_image_paths, should_keep_image


class PPTXParser(BaseParser):
    def __init__(self, logger=None, output_images_dir: str = os.path.join("output", "default", "images")):
        super().__init__(logger)
        self.output_images_dir = output_images_dir

    def parse(self, file_path: str) -> Dict[str, Any]:
        if not self.validate_file(file_path):
            self.log_error(f"File invalid: {file_path}")
            return {"success": False, "error": "invalid file"}

        try:
            self.log_info(f"Start parsing PPTX: {file_path}")
            prs = Presentation(file_path)
            filename = os.path.splitext(os.path.basename(file_path))[0]

            pages = []
            for slide_num in range(len(prs.slides)):
                self.log_info(f"Parsing slide {slide_num + 1}/{len(prs.slides)}")
                pages.append(self.extract_page_content(file_path, slide_num))

            image_mapping = self.extract_all_images_via_zip(file_path)

            metadata = self.get_metadata(file_path)
            metadata.update({"total_slides": len(prs.slides), "image_mapping": image_mapping})

            result = {
                "success": True,
                "filename": filename,
                "pages": pages,
                "metadata": metadata,
            }
            self.log_info(f"PPTX parsed: {filename}")
            return result
        except Exception as e:
            self.log_error(f"PPTX parse failed: {e}")
            return {"success": False, "error": str(e)}

    def get_page_count(self, file_path: str) -> int:
        try:
            prs = Presentation(file_path)
            return len(prs.slides)
        except Exception as e:
            self.log_error(f"Get slide count failed: {e}")
            return 0

    def extract_page_content(self, file_path: str, page_num: int) -> Dict[str, Any]:
        try:
            prs = Presentation(file_path)
            slide = prs.slides[page_num]

            text_parts: List[str] = []
            tables = []

            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            text_parts.append(run.text)

                if shape.shape_type == MSO_SHAPE_TYPE.TABLE:
                    tables.append(
                        {
                            "shape": shape,
                            "position": (shape.left, shape.top, shape.width, shape.height),
                        }
                    )

            return {
                "page_num": page_num,
                "text": "\n".join(text_parts),
                "images": [],
                "tables": tables,
                "raw_slide": slide,
            }
        except Exception as e:
            self.log_error(f"Extract slide content failed (slide {page_num}): {e}")
            return {
                "page_num": page_num,
                "text": "",
                "images": [],
                "tables": [],
                "raw_slide": None,
                "error": str(e),
            }

    def extract_all_images_via_zip(self, file_path: str) -> Dict[int, List[str]]:
        self.log_info("Extract images via zip")

        filename = os.path.splitext(os.path.basename(file_path))[0]
        slide_image_mapping: Dict[int, List[str]] = {}

        with tempfile.TemporaryDirectory() as temp_dir:
            try:
                with zipfile.ZipFile(file_path, "r") as zip_ref:
                    zip_ref.extractall(temp_dir)

                media_dir = os.path.join(temp_dir, "ppt", "media")
                slides_dir = os.path.join(temp_dir, "ppt", "slides")
                rels_dir = os.path.join(temp_dir, "ppt", "slides", "_rels")

                if not os.path.exists(media_dir):
                    self.log_info("No media directory found")
                    return slide_image_mapping

                if os.path.exists(slides_dir):
                    for slide_file in os.listdir(slides_dir):
                        if not (slide_file.startswith("slide") and slide_file.endswith(".xml")):
                            continue

                        slide_num = self._extract_slide_number(slide_file)
                        if slide_num is None:
                            continue

                        rels_file = os.path.join(rels_dir, f"{slide_file}.rels")
                        if not os.path.exists(rels_file):
                            continue

                        image_files = self._get_slide_images(rels_file, media_dir)
                        output_images: List[str] = []

                        for idx, img_file in enumerate(image_files):
                            file_ext = os.path.splitext(img_file)[1].lower()
                            output_path = os.path.join(
                                self.output_images_dir,
                                f"{filename}_slide_{slide_num}_img_{idx+1}{file_ext}",
                            )
                            os.makedirs(os.path.dirname(output_path), exist_ok=True)

                            shutil.copy2(img_file, output_path)

                            if file_ext in [".emf", ".wmf", ".svg"]:
                                converted_path = self._convert_vector_to_png(output_path)
                                if converted_path:
                                    output_path = converted_path
                                else:
                                    self.log_warning(f"Skip unconvertible vector image: {output_path}")
                                    try:
                                        os.remove(output_path)
                                    except Exception:
                                        pass
                                    continue

                            keep, reason, stats = should_keep_image(output_path)
                            if not keep:
                                try:
                                    os.remove(output_path)
                                except Exception:
                                    pass
                                self.log_info(
                                    f"Skip low-quality PPTX image: {os.path.basename(output_path)} | reason={reason} | stats={stats}"
                                )
                                continue

                            output_images.append(output_path)

                        if output_images:
                            slide_image_mapping[slide_num] = output_images

                self._remove_repeated_images(slide_image_mapping)
                self.log_info(
                    f"Extracted {sum(len(imgs) for imgs in slide_image_mapping.values())} images"
                )
            except Exception as e:
                self.log_error(f"ZIP image extraction failed: {e}")

        return slide_image_mapping

    def _remove_repeated_images(self, slide_image_mapping: Dict[int, List[str]], min_count: int = 3):
        all_images = []
        for images in slide_image_mapping.values():
            all_images.extend(images)

        repeated = repeated_image_paths(all_images, min_count=min_count)
        if not repeated:
            return

        for slide_num, images in list(slide_image_mapping.items()):
            kept = []
            for image_path in images:
                if image_path in repeated:
                    try:
                        os.remove(image_path)
                    except Exception:
                        pass
                    self.log_info(
                        f"Skip repeated PPTX logo/watermark: {os.path.basename(image_path)} | signature={repeated[image_path]}"
                    )
                else:
                    kept.append(image_path)
            if kept:
                slide_image_mapping[slide_num] = kept
            else:
                slide_image_mapping.pop(slide_num, None)

    def _extract_slide_number(self, slide_filename: str) -> int:
        match = re.search(r"slide(\d+)", slide_filename)
        if match:
            return int(match.group(1))
        return None

    def _get_slide_images(self, rels_file: str, media_dir: str) -> List[str]:
        images = []
        try:
            tree = ET.parse(rels_file)
            root = tree.getroot()

            for rel in root.findall(
                ".//{http://schemas.openxmlformats.org/package/2006/relationships}Relationship"
            ):
                rel_type = rel.get("Type", "")
                if "image" in rel_type.lower():
                    target = rel.get("Target", "")
                    img_filename = os.path.basename(target)
                    img_path = os.path.join(media_dir, img_filename)
                    if os.path.exists(img_path):
                        images.append(img_path)
        except Exception as e:
            self.log_error(f"Parse rels failed: {e}")

        return images

    def _convert_vector_to_png(self, image_path: str) -> str:
        file_ext = os.path.splitext(image_path)[1].lower()
        if file_ext not in [".emf", ".wmf", ".svg"]:
            return image_path

        png_path = image_path.rsplit(".", 1)[0] + "_converted.png"
        self.log_info(f"Convert vector image: {file_ext} -> PNG")

        try:
            import subprocess

            try:
                subprocess.run(["magick", "-version"], capture_output=True, check=True)
                has_imagemagick = True
            except Exception:
                has_imagemagick = False

            if has_imagemagick:
                cmd = [
                    "magick",
                    "convert",
                    image_path,
                    "-background",
                    "white",
                    "-alpha",
                    "remove",
                    png_path,
                ]
                result = subprocess.run(cmd, capture_output=True, text=True)
                if result.returncode == 0 and os.path.exists(png_path):
                    return png_path
        except Exception as e:
            self.log_warning(f"ImageMagick conversion failed: {e}")

        try:
            from wand.image import Image as WandImage

            with WandImage(filename=image_path, resolution=300) as img:
                img.background_color = "white"
                img.alpha_channel = "remove"
                img.format = "png"
                img.save(filename=png_path)

            if os.path.exists(png_path):
                return png_path
        except ImportError:
            self.log_info("wand not installed, vector image conversion unavailable")
        except Exception as e:
            self.log_warning(f"Wand conversion failed: {e}")

        self.log_warning(f"Skip vector image after conversion failure: {image_path}")
        return None

    def extract_slide_text_only(self, file_path: str, slide_num: int) -> str:
        try:
            prs = Presentation(file_path)
            slide = prs.slides[slide_num]

            text_parts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        for run in paragraph.runs:
                            text_parts.append(run.text)

            return "\n".join(text_parts)
        except Exception as e:
            self.log_error(f"Extract slide text failed (slide {slide_num}): {e}")
            return ""



